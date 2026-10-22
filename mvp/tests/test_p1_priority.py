"""Priority regressions use synthetic tenants and the isolated QA database only."""
import time
import pytest
from sqlalchemy import select
from server.turn_plan import direction_state, build_turn_plan
from server.dialogue_policy import dialogue_context
from mvp.document_parser import normalize, chunks
from mvp.db import transaction, rows
from mvp.models import Knowledge, KBChunk, KBDocument, Conversation, Message, Channel, ChannelDelivery
from test_conversation import send
from test_document_knowledge import ready


def history(*texts):
    return [{'role':'user','content':text} for text in texts]


@pytest.mark.parametrize('question',['你们做过官网吗？','能做官网吗？','官网有哪些案例？'])
def test_service_question_does_not_replace_choice(question):
    state=direction_state(history('先做管理系统',question))
    assert state['status']=='selected' and state['candidates']==['管理系统']


@pytest.mark.parametrize('text,expected',[
    ('官网不做了，先做小程序',['小程序']),
    ('暂时不要做官网了',[]),
    ('不做官网，改做管理系统',['管理系统']),
    ('不是官网',[]),
])
def test_direction_cancellation(text,expected):
    assert direction_state(history('先做官网',text))['candidates']==expected


def test_project_restart_resets_all_context():
    h=history('先做官网，预算8万，下个月上线','换个项目，重新开始')
    assert build_turn_plan(h)['budget_status']=='unknown'
    ctx=dialogue_context(h)
    assert ctx['topics']==[] and ctx['user_requirements']=={}


def test_project_restart_is_not_a_hypothetical_action():
    h=history('先做官网，预算8万','如果换个项目，重新开始会怎样？')
    assert direction_state(h)['candidates']==['企业官网']
    assert build_turn_plan(h)['budget_status']=='set'


@pytest.mark.parametrize('question',['不想重新开始','你们新项目报价有哪些？','如果换个项目，预算3万元呢？','如果换个项目，重新开始，预算3万'])
def test_restart_negation_or_question_is_not_a_new_project(question):
    ctx=dialogue_context(history('先做官网，预算8万',question))
    assert '8万' in ctx['user_requirements']['budget']


def test_restart_inside_a_message_drops_only_prior_project_facts():
    h=history('先做官网，预算8万','之前预算8万。换个项目，想做小程序，预算3万')
    assert direction_state(h)['candidates']==['小程序']
    assert '3万' in dialogue_context(h)['user_requirements']['budget']


def test_existing_stale_profile_is_repaired_on_next_message(make_tenant,make_visitor):
    t=make_tenant();v,cid=make_visitor(t)
    send(t,v,cid,'先做官网，预算8万，下个月上线','old')
    send(t,v,cid,'换个项目，重新开始','new')
    with transaction(t.id) as s:
        c=s.get(Conversation,(t.id,cid));c.profile={'budget':{'value':80000,'display':'¥80,000','evidence':'旧项目','confirmed':True}}
    send(t,v,cid,'先做小程序','current')
    profile=v.get('/api/public/conversations/'+cid).json()['profile']
    assert profile['need']['value']=='小程序' and 'budget' not in profile


def test_docx_heading_levels_and_table_conditions_survive(tmp_path):
    from docx import Document
    p=tmp_path/'policy.docx';d=Document()
    d.add_heading('设备政策',level=1);d.add_heading('仅限续约客户',level=2);d.add_heading('参数清单',level=3)
    table=d.add_table(rows=2,cols=2)
    table.cell(0,0).text='项目';table.cell(0,1).text='内容'
    table.cell(1,0).text='维护费';table.cell(1,1).text='800元每年'
    d.save(p)
    part=chunks(normalize(p,p.name),max_chars=360)[0]
    assert '仅限续约客户' in part['content'] and '参数清单' in part['content']
    assert part['location']['heading_path']==['设备政策','仅限续约客户','参数清单']


def test_project_ack_does_not_need_a_decorative_company_citation(monkeypatch):
    import json
    from server import model_gateway
    monkeypatch.setattr(model_gateway,'complete',lambda *_:json.dumps({'answer':'好的，按新项目重新梳理。您这次想解决什么问题呢？','knowledge_ids':[],'covered_tasks':['project_restart'],'case_ids':[]},ensure_ascii=False))
    answer,used=model_gateway.grounded_answer(dict(model_gateway.DEFAULT),'合成企业',history('换个项目，重新开始'),[{'id':'unrelated','title':'企业介绍','content':'提供官网咨询。','version':1}])
    assert '新项目' in answer and not used


def test_profile_reply_and_handoff_use_the_same_project(make_tenant,make_visitor,make_employee):
    t=make_tenant();v,cid=make_visitor(t);a,_=make_employee(t)
    send(t,v,cid,'先做官网，预算8万，下个月上线','initial')
    reply=send(t,v,cid,'换个项目，重新开始','restart')
    assert '新项目' in reply['text'] or '重新' in reply['text']
    detail=v.get('/api/public/conversations/'+cid).json()
    assert not {'need','budget','timeline'}.intersection(detail['profile'])
    assert '8万' not in reply['text'] and '下个月' not in reply['text']
    assert not a.get('/api/conversations/'+cid).json()['handoff_summary']['confirmed_needs']
    send(t,v,cid,'我想做管理系统，预算3万元','new')
    assert v.get('/api/public/conversations/'+cid).json()['profile']['budget']['value']==30000


def test_profile_question_preserves_previous_direction(make_tenant,make_visitor):
    t=make_tenant();v,cid=make_visitor(t)
    send(t,v,cid,'先做管理系统，预算8万','select')
    send(t,v,cid,'你们做过官网吗？','question')
    detail=v.get('/api/public/conversations/'+cid).json()
    assert detail['profile']['need']['value']=='管理系统'
    assert detail['profile']['budget']['value']==80000
    cancelled=send(t,v,cid,'管理系统暂时不做了','cancel')
    assert any(x in cancelled['text'] for x in ['暂停','不做','不再','取消'])
    assert 'need' not in v.get('/api/public/conversations/'+cid).json()['profile']


@pytest.mark.parametrize('blank',['\n','\n\n'])
def test_parent_conditions_survive_nested_markdown(tmp_path,blank):
    p=tmp_path/'policy.md'
    p.write_text(blank.join(['# 官网优惠政策','## 仅限续约老客户','### 企业官网价格','官网维护费800元每年。','## 新客户','### 企业官网价格','新客户维护费1200元每年。']),encoding='utf-8')
    pieces=chunks(normalize(p,p.name),max_chars=360)
    old=next(c for c in pieces if '800元' in c['content'])
    new=next(c for c in pieces if '1200元' in c['content'])
    assert '仅限续约老客户' in old['content']
    assert '仅限续约老客户' not in new['content'] and '新客户' in new['content']


@pytest.mark.parametrize('price,content',[
    (99,'企业官网服务公开价格99999元。'),
    (99,'企业官网服务支持99人使用。'),
    (99,'企业官网服务退款99元，服务报价99999元。'),
])
def test_public_price_requires_the_actual_price(make_tenant,make_employee,price,content):
    t=make_tenant();a,_=make_employee(t)
    with transaction(t.id) as s:
        s.add(Knowledge(tenant_id=t.id,id='price',family='price',title='企业官网服务价格',content=content,
            status='published',risk='high',author_id='author',reviewer_id='reviewer'))
    response=a.put('/api/offers/new',json={'name':'企业官网服务','description':'测试服务','tags':['企业官网'],'price':price,'knowledge_id':'price','active':True})
    assert response.status_code==409


@pytest.mark.parametrize('content',['企业官网服务公开价格99元。','企业官网服务公开价格0.0099万元。'])
def test_equivalent_public_price_is_accepted(make_tenant,make_employee,content):
    t=make_tenant();a,_=make_employee(t)
    with transaction(t.id) as s:
        s.add(Knowledge(tenant_id=t.id,id='price',family='price',title='企业官网服务价格',content=content,
            status='published',risk='high',author_id='author',reviewer_id='reviewer'))
    assert a.put('/api/offers/new',json={'name':'企业官网服务','description':'测试服务','tags':['企业官网'],'price':99,'knowledge_id':'price','active':True}).status_code==200


@pytest.mark.parametrize('content',[
    '企业官网服务价格50至99元。',
    '企业官网服务价格99元起。',
    '企业官网服务价格99元每年。',
    '仅限老客户，企业官网服务价格99元。',
    '如果采购，企业官网服务价格99元。',
    '企业官网服务价格999元，管理系统价格99元。',
    '企业官网服务价格999元和管理系统价格99元。',
    '企业官网服务价格99欧元。',
])
def test_price_conditions_and_objects_do_not_become_flat_prices(content):
    from mvp.offer_evidence import supports_public_price
    assert not supports_public_price({'content':content},'企业官网服务',99)


def test_empty_service_name_cannot_authorize_a_price():
    from mvp.offer_evidence import supports_public_price
    assert not supports_public_price({'content':'企业官网服务价格99元。'},'   ',99)


def test_legacy_invalid_price_is_hidden_and_can_be_disabled(make_tenant,make_visitor,make_employee):
    from mvp.models import Offer
    t=make_tenant();v,cid=make_visitor(t);a,_=make_employee(t)
    with transaction(t.id) as s:
        s.add(Knowledge(tenant_id=t.id,id='price',family='price',title='企业官网服务价格',content='企业官网服务价格99999元。',status='published',risk='high',author_id='author',reviewer_id='reviewer'))
        s.add(Offer(tenant_id=t.id,id='old',name='企业官网服务',description='旧服务',tags=['企业官网'],price=99,active=True,knowledge_id='price'))
    send(t,v,cid,'先做官网，预算10万','select')
    assert not v.get('/api/public/conversations/'+cid).json()['recommendations']
    assert a.put('/api/offers/old',json={'name':'企业官网服务','description':'旧服务','tags':['企业官网'],'price':99,'knowledge_id':'price','active':False}).status_code==200


def test_correct_price_disappears_after_source_change(make_tenant,make_visitor,make_employee):
    t=make_tenant();a,_=make_employee(t);v,cid=make_visitor(t)
    with transaction(t.id) as s:
        s.add(Knowledge(tenant_id=t.id,id='price',family='price',title='企业官网服务价格',content='企业官网服务价格99元。',status='published',risk='high',author_id='author',reviewer_id='reviewer'))
    assert a.put('/api/offers/new',json={'name':'企业官网服务','description':'测试服务','tags':['企业官网'],'price':99,'knowledge_id':'price','active':True}).status_code==200
    send(t,v,cid,'先做官网，预算1万','select')
    assert v.get('/api/public/conversations/'+cid).json()['recommendations'][0]['price']==99
    with transaction(t.id) as s:s.get(Knowledge,(t.id,'price')).status='withdrawn'
    assert not v.get('/api/public/conversations/'+cid).json()['recommendations']


def test_document_v2_can_be_delivered_and_withdrawal_still_blocks(make_tenant,make_employee,monkeypatch):
    from test_channels import create, frame
    from mvp.channel_messages import ingest, check_delivery
    from mvp.channel_transport import PlatformError
    from server.retrieval import citation
    from mvp.knowledge_retrieval import PostgresScan
    t=make_tenant();a,_=make_employee(t);ch=create(a,monkeypatch,'dingtalk')
    first,_=ready(t,a,content='# 官网服务\n\n支持企业官网内容展示。',visibility='public')
    assert a.post('/api/kb/versions/'+first['version_id']+'/publish',json={}).status_code==200
    col=a.get('/api/kb/collections').json()[0]['id']
    uploaded=a.post('/api/kb/documents/upload',data={'collection_id':col,'document_id':first['document_id'],'visibility':'public'},files={'file':('产品.md','# 官网服务\n\n支持企业官网和产品资料展示。'.encode())})
    assert uploaded.status_code==200,uploaded.text
    from mvp.knowledge_worker import lease
    from mvp.knowledge_documents import process
    process(t.id,lease(t.id));second=uploaded.json()
    assert a.post('/api/kb/versions/'+second['version_id']+'/publish',json={}).status_code==200
    cid,_=ingest(t.id,ch['id'],ch['version'],frame(text='企业官网资料'))
    with transaction(t.id) as s:
        record=next(r for r,_,_ in PostgresScan(s,t.id,ch['id']).scan() if r.get('document_id')==first['document_id'])
        assert record['version']==2
        from mvp.service import add_message
        c=s.get(Conversation,(t.id,cid));m=add_message(s,c,'ai','支持产品资料展示。',citations=[citation(record)])
        delivery=s.scalar(rows(s,ChannelDelivery,t.id).where(ChannelDelivery.message_id==m.id))
        check_delivery(s,delivery)
        # Historic references without an evidence snapshot remain supported,
        # but still require canonical active content and channel authorization.
        m.citations=[{k:v for k,v in m.citations[0].items() if k!='evidence'}]
        check_delivery(s,delivery)
        doc=s.get(KBDocument,(t.id,first['document_id']))
        doc.channel_ids=['another-channel'];s.flush()
        with pytest.raises(PlatformError):check_delivery(s,delivery)
        doc.channel_ids=[];s.flush()
        knowledge=s.get(Knowledge,(t.id,record['id']));original=knowledge.content
        knowledge.content='被改写的失效内容';s.flush()
        with pytest.raises(PlatformError):check_delivery(s,delivery)
        knowledge.content=original;s.flush()
        check_delivery(s,delivery)
        doc.active_version=None
        s.flush()
        with pytest.raises(PlatformError):check_delivery(s,delivery)
