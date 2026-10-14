import io
import json
import pytest
from sqlalchemy import select,text
from mvp.db import transaction,rows
from mvp.models import *
from mvp.knowledge_worker import lease
from mvp.knowledge_documents import process
from mvp.document_parser import normalize,chunks,ParseError
from mvp.knowledge_engine import hybrid,cosine

def test_delete_removes_original_index_and_public_source(make_tenant,make_employee):
    from mvp.knowledge_documents import storage_path
    t=make_tenant();a,_=make_employee(t);d,_=ready(t,a,visibility='public')
    a.post('/api/kb/versions/'+d['version_id']+'/publish',json={})
    with transaction(t.id) as s:path=storage_path(s.get(KBDocumentVersion,(t.id,d['version_id'])).storage_key)
    assert path.exists()
    assert a.delete('/api/kb/documents/'+d['document_id']).status_code==200
    assert not path.exists()
    assert a.get('/api/kb/versions/'+d['version_id']+'/original').status_code==404
    with transaction(t.id) as s:assert not list(s.scalars(rows(s,KBChunk,t.id)))

def test_document_chunks_cannot_be_edited_through_legacy_api(make_tenant,make_employee):
    t=make_tenant();a,_=make_employee(t);d,_=ready(t,a,visibility='public')
    a.post('/api/kb/versions/'+d['version_id']+'/publish',json={})
    with transaction(t.id) as s:c=s.scalar(rows(s,KBChunk,t.id))
    assert a.post('/api/knowledge',json={'family':c.id,'title':'改写','content':'这是通过旧接口直接修改的知识内容。'}).status_code==409
    assert a.post('/api/knowledge/'+c.id+'/withdraw',json={'version':1}).status_code==409

def test_generic_support_profile_is_not_website_sales(make_tenant,make_employee,make_visitor):
    from test_conversation import send
    t=make_tenant();a,_=make_employee(t)
    with transaction(t.id) as s:s.get(CompanySettings,(t.id,'primary')).service_mode='general_support'
    v,cid=make_visitor(t);reply=send(t,v,cid,'电视型号XYZ如何连接网络？')
    with transaction(t.id) as s:p=s.get(Conversation,(t.id,cid)).profile
    assert p['need']['label']=='咨询内容' and '建站' not in json.dumps(p,ensure_ascii=False)
    assert '官网建设' not in reply['text']

def test_scanned_pdf_is_reported_as_unreadable(tmp_path):
    from pypdf import PdfWriter
    w=PdfWriter();w.add_blank_page(width=300,height=300);path=tmp_path/'scan.pdf'
    with path.open('wb') as f:w.write(f)
    with pytest.raises(ParseError,match='ocr_required'):normalize(path,path.name)

def test_each_section_retains_parent_product_identity(tmp_path):
    p=tmp_path/'manual.md';p.write_text('# 产品ZX900\n\n## 网络设置\n\n支持Wi-Fi。\n\n## 功耗\n\n功率200W。','utf-8')
    parts=chunks(normalize(p,p.name),max_chars=360)
    assert len(parts)==2 and all('ZX900' in x['title'] and 'ZX900' in x['content'] for x in parts)

def test_batch_publish_preserves_review_and_tenant_checks(make_tenant,make_employee):
    t=make_tenant();a,_=make_employee(t);one,_=ready(t,a,visibility='public')
    two,_=ready(t,a,content='产品ABC价格100元。',name='报价.md',visibility='public')
    r=a.post('/api/kb/publish-batch',json={'version_ids':[one['version_id'],two['version_id'],'unknown']})
    assert r.status_code==200
    results=r.json()['results'];assert [x['status'] for x in results]==['published','blocked','blocked']
    assert results[1]['code']==403 and results[2]['code']==404

def test_initial_admin_invite_reissue_never_resets_existing_account():
    from fastapi.testclient import TestClient
    from mvp.app import app
    c=TestClient(app);token=c.get('/api/platform/local-token').json()['token'];c.headers['X-Operator-Token']=token
    slug='invite-'+uid()[:8]
    r=c.post('/api/platform/tenants',json={'slug':slug,'name':'合成试验企业','username':'admin1','admin_name':'试验管理员'})
    assert r.status_code==200;old=r.json()['invite_url'].split('#')[1]
    r=c.post('/api/platform/tenants/'+slug+'/admin-invitation');assert r.status_code==200;new=r.json()['invite_url'].split('#')[1]
    assert c.post('/api/invitations/accept',json={'tenant':slug,'token':old,'password':'synthetic-test-password-only'}).status_code==403
    assert c.post('/api/invitations/accept',json={'tenant':slug,'token':new,'password':'synthetic-test-password-only'}).status_code==200
    assert c.post('/api/platform/tenants/'+slug+'/admin-invitation').status_code==409

def test_generic_model_off_can_quote_correct_field_without_cross_product_merge():
    from mvp.answering import support_fallback
    records=[{'id':'identity','title':'产品ABC · 身份','content':'ABC是测试产品。','document_id':'a'},
             {'id':'different','title':'产品DEF · 参数','content':'| 产品类型 | 波轮洗衣机 |','document_id':'b'},
             {'id':'correct','title':'产品ABC · 参数','content':'| 产品类型 | 全自动滚筒式洗衣机 |','document_id':'a'}]
    answer,used=support_fallback(records,'ABC是什么类型的洗衣机？')
    assert '全自动滚筒式洗衣机' in answer and '波轮' not in answer and used[0]['id']=='correct'

def test_date_pair_cannot_be_spliced_from_two_sources():
    from server.answer_policy import date_ranges_supported
    evidence=[{'content':'活动时间11月1日至11月7日。'},{'content':'活动时间11月2日至11月8日。'}]
    assert not date_ranges_supported('活动时间为11月1日至11月8日。',evidence)
    assert date_ranges_supported('此来源的活动时间11月2日至11月8日。',evidence)

def test_conflicting_same_title_date_requires_replacement_version(make_tenant,make_employee):
    t=make_tenant();a,_=make_employee(t)
    one,_=ready(t,a,content='# 活动时间\n\n活动时间11月1日至11月7日。',name='活动.md',visibility='public')
    assert a.post('/api/kb/versions/'+one['version_id']+'/publish',json={}).status_code==200
    two,_=ready(t,a,content='# 活动时间\n\n活动时间11月2日至11月8日。',name='活动.md',visibility='public')
    assert a.post('/api/kb/versions/'+two['version_id']+'/publish',json={}).status_code==409

def test_explicit_unknown_model_never_matches_similar_digits():
    from mvp.knowledge_engine import scope_by_model
    records=[{'id':'wrong','title':'BD/BC-100E1SHMKD','content':'这是100L冷柜。'}]
    result,diagnostic=hybrid('星叶T100',records,{}, {'embedding_mode':'keyword','rerank_mode':'off'}, {})
    assert result==[] and diagnostic['model_scope']=='unavailable'
    assert scope_by_model('T1000',records+[{'title':'T100','content':'净化器'}])[0]==[]
    assert scope_by_model('星叶t100适用多大面积？',records+[{'title':'星叶T100 · 参数','content':'面积8至12平方米'}])[0][0]['title'].startswith('星叶')
    assert scope_by_model('功率18W和WiFi6是什么意思',records)[0]==records

def test_private_upload_scope_can_change_only_after_withdraw_and_republish(make_tenant,make_employee,make_visitor):
    from test_conversation import send
    t=make_tenant();a,_=make_employee(t);v,cid=make_visitor(t)
    d,_=ready(t,a,content='星叶T100是空气净化器，适用面积8至12平方米。')
    a.post('/api/kb/versions/'+d['version_id']+'/publish',json={})
    path='/api/kb/documents/'+d['document_id']
    assert a.put(path+'/scope',json={'visibility':'public'}).status_code==409
    assert not send(t,v,cid,'星叶T100',client_id='private')['citations']
    a.post(path+'/withdraw');assert a.put(path+'/scope',json={'visibility':'public'}).status_code==200
    assert not send(t,v,cid,'星叶T100',client_id='unpublished')['citations']
    assert a.post('/api/kb/versions/'+d['version_id']+'/publish',json={}).status_code==200
    reply=send(t,v,cid,'星叶T100',client_id='published')
    assert reply['citations'] and 'T100' in reply['text']

def test_scope_change_rejects_foreign_documents_channels_and_agent(make_tenant,make_employee):
    t=make_tenant();b=make_tenant();a,_=make_employee(t);other,_=make_employee(b);agent,_=make_employee(t,roles=['agent'])
    d,_=ready(t,a);path='/api/kb/documents/'+d['document_id']+'/scope'
    assert agent.put(path,json={'visibility':'public'}).status_code==403
    assert other.put(path,json={'visibility':'public'}).status_code==404
    assert a.put(path,json={'visibility':'public','channel_ids':['foreign']}).status_code==422

def test_plain_product_fragment_retains_later_facts_and_does_not_invent_price():
    from mvp.answering import support_fallback
    r={'title':'星叶T100 · 身份','document_id':'test','content':'星叶T100 · 身份\n测试品牌：星叶。型号T100。适用面积8至12平方米。'}
    answer,used=support_fallback([r],'T100适用多大面积？')
    assert '8至12平方米' in answer and used==[r]
    answer,used=support_fallback([r],'星叶T100多少钱？')
    assert '没有可核对的售价' in answer and '100元' not in answer

def ready(t,client,content='企业产品ABC是工业设备。用于生产线监测，额定功率200W。',visibility='internal',name='产品.md',data=None):
    with transaction(t.id) as s:
        if not s.get(KBEngineConfig,(t.id,'primary')):s.add(KBEngineConfig(tenant_id=t.id,id='primary',config={'embedding_mode':'keyword','rerank_mode':'off'}))
    collection=client.get('/api/kb/collections').json()[0]['id']
    response=client.post('/api/kb/documents/upload',data={'collection_id':collection,'visibility':visibility},files={'file':(name,data or content.encode('utf-8'))})
    assert response.status_code==200,response.text
    result=response.json()
    if not result['duplicate']:process(t.id,lease(t.id))
    return result,collection

def test_uploaded_doc_is_draft_internal_and_source_preserving(make_tenant,make_employee):
    t=make_tenant();a,_=make_employee(t);d,_=ready(t,a)
    detail=a.get('/api/kb/documents/'+d['document_id']).json()
    assert detail['visibility']=='internal' and detail['versions'][0]['status']=='ready'
    standard=a.get('/api/kb/versions/'+d['version_id']+'/standard')
    assert standard.status_code==200 and standard.json()['schema']=='salescare.document.v1'
    assert standard.json()['blocks'][0]['location']['paragraph']==1
    assert a.get('/api/kb/versions/'+d['version_id']+'/original').content==standard.json()['blocks'][0]['text'].encode('utf-8')

def test_duplicate_upload_does_not_create_job_or_version(make_tenant,make_employee):
    t=make_tenant();a,_=make_employee(t);one,_=ready(t,a);two,_=ready(t,a)
    assert two['duplicate'] and one['version_id']==two['version_id']
    with transaction(t.id) as s:assert len(list(s.scalars(rows(s,KBJob,t.id))))==1

def test_tenant_object_and_storage_access_are_isolated(make_tenant,make_employee):
    a=make_tenant();b=make_tenant();ca,_=make_employee(a);cb,_=make_employee(b);d,_=ready(a,ca)
    assert cb.get('/api/kb/documents/'+d['document_id']).status_code==404
    assert cb.get('/api/kb/versions/'+d['version_id']+'/original').status_code==404
    assert cb.get('/api/kb/versions/'+d['version_id']+'/preview').status_code==404
    assert cb.post('/api/kb/versions/'+d['version_id']+'/publish',json={}).status_code==404

def test_agent_cannot_upload_configure_or_read_internal_documents(make_tenant,make_employee):
    t=make_tenant();a,_=make_employee(t);agent,_=make_employee(t,roles=['agent']);d,col=ready(t,a)
    assert agent.get('/api/kb/documents').status_code==403
    assert agent.get('/api/kb/engine').status_code==403
    assert agent.post('/api/kb/documents/upload',data={'collection_id':col},files={'file':('x.txt',b'test text')}).status_code==403

def test_editor_and_publisher_can_load_nonsecret_channels(make_tenant,make_employee):
    t=make_tenant();e,_=make_employee(t,roles=['editor']);p,_=make_employee(t,roles=['publisher'])
    assert e.get('/api/kb/channels').status_code==200 and p.get('/api/kb/channels').status_code==200
    assert e.get('/api/kb/engine').status_code==403

def test_internal_publication_never_enters_consumer_knowledge(make_tenant,make_employee,make_visitor):
    from test_conversation import send
    t=make_tenant();a,_=make_employee(t);v,cid=make_visitor(t);d,_=ready(t,a,content='保密型号XYZ，内部采购成本7777。')
    assert a.post('/api/kb/versions/'+d['version_id']+'/publish',json={}).status_code==200
    assert '7777' not in send(t,v,cid,'XYZ是什么？')['text']
    with transaction(t.id) as s:c=s.scalar(rows(s,KBChunk,t.id))
    assert v.get('/api/public/document-source/'+c.id).status_code==404

def test_public_document_can_answer_and_withdraw_removes_source(make_tenant,make_employee,make_visitor):
    from test_conversation import send
    t=make_tenant();a,_=make_employee(t);v,cid=make_visitor(t);d,_=ready(t,a,content='产品XYZ用于工业检测，支持温度采集。',visibility='public')
    assert a.post('/api/kb/versions/'+d['version_id']+'/publish',json={}).json()['status']=='published'
    reply=send(t,v,cid,'XYZ工业检测产品是什么？')
    assert 'XYZ' in reply['text'] and reply['citations']
    key=reply['citations'][0]['id']
    assert v.get('/api/public/document-source/'+key).status_code==200
    assert a.post('/api/kb/documents/'+d['document_id']+'/withdraw',json={}).status_code==200
    assert v.get('/api/public/document-source/'+key).status_code==404

def test_new_candidate_does_not_change_active_version_until_publish(make_tenant,make_employee):
    t=make_tenant();a,_=make_employee(t);d,col=ready(t,a,content='产品ABC参数为200W。',visibility='public')
    a.post('/api/kb/versions/'+d['version_id']+'/publish',json={})
    newer=a.post('/api/kb/documents/upload',data={'collection_id':col,'document_id':d['document_id']},files={'file':('产品.md','产品ABC参数修订为300W。'.encode())}).json()
    process(t.id,lease(t.id))
    assert a.get('/api/kb/documents/'+d['document_id']).json()['active_version']==d['version_id']
    a.post('/api/kb/versions/'+newer['version_id']+'/publish',json={})
    with transaction(t.id) as s:
        old=list(s.scalars(rows(s,KBChunk,t.id).where(KBChunk.version_id==d['version_id'])))
        assert all(s.get(Knowledge,(t.id,c.id)).status=='superseded' for c in old)

def test_high_risk_document_cannot_be_self_approved(make_tenant,make_employee):
    t=make_tenant();a,_=make_employee(t);d,_=ready(t,a,content='测试产品价格100元，退款需核实。',visibility='public')
    assert a.post('/api/kb/versions/'+d['version_id']+'/publish',json={}).status_code==403
    publisher,_=make_employee(t,roles=['publisher'])
    assert publisher.post('/api/kb/versions/'+d['version_id']+'/publish',json={}).status_code==200

def test_failed_file_can_retry_without_modifying_published_data(make_tenant,make_employee):
    t=make_tenant();a,_=make_employee(t);d,_=ready(t,a,name='empty.md',data=b'   ')
    assert a.get('/api/kb/documents/'+d['document_id']).json()['versions'][0]['status']=='failed'
    assert a.post('/api/kb/versions/'+d['version_id']+'/retry',json={}).status_code==200
    with transaction(t.id) as s:assert s.get(KBDocumentVersion,(t.id,d['version_id'])).status=='queued'

def test_parser_keeps_table_headers_units_and_rows(make_tenant,make_employee):
    from openpyxl import Workbook
    book=Workbook();book.active.title='产品参数';book.active.append(['型号','功率W']);book.active.append(['ABC',200]);book.active.append(['XYZ',300]);buf=io.BytesIO();book.save(buf)
    t=make_tenant();a,_=make_employee(t);d,_=ready(t,a,name='产品.xlsx',data=buf.getvalue())
    result=a.get('/api/kb/versions/'+d['version_id']+'/preview').json()
    table=result['blocks'][0];assert table['rows'][1]==['ABC','200'] and table['header']==['型号','功率W']
    with transaction(t.id) as s:
        c=s.scalar(rows(s,KBChunk,t.id));assert '功率W' in c.content and c.location['sheet']=='产品参数'

def test_docx_and_pptx_standardize_without_manual_markdown(make_tenant,make_employee):
    from docx import Document
    from pptx import Presentation
    word=Document();word.add_heading('产品介绍',0);word.add_paragraph('产品ABC用于监控。');buf=io.BytesIO();word.save(buf)
    t=make_tenant();a,_=make_employee(t);d,_=ready(t,a,name='资料.docx',data=buf.getvalue())
    assert a.get('/api/kb/versions/'+d['version_id']+'/preview').json()['blocks']
    presentation=Presentation();slide=presentation.slides.add_slide(presentation.slide_layouts[1]);slide.shapes.title.text='公司产品';slide.placeholders[1].text='ABC产品用于监测。';buf=io.BytesIO();presentation.save(buf)
    d,_=ready(t,a,name='资料.pptx',data=buf.getvalue())
    assert a.get('/api/kb/versions/'+d['version_id']+'/preview').json()['blocks'][0]['location']['slide']==1

def test_preview_sees_drafts_but_consumer_does_not(make_tenant,make_employee):
    t=make_tenant();a,_=make_employee(t);ready(t,a,content='产品XYZ支持温度监测。')
    result=a.post('/api/kb/search',json={'query':'XYZ温度监测'}).json()
    assert result['results'] and 'XYZ' in result['results'][0]['content']
    with transaction(t.id) as s:assert not s.scalar(rows(s,Knowledge,t.id).where(Knowledge.category=='文档知识'))

def test_embedding_fallback_and_rrf_keep_precise_identifiers(monkeypatch):
    from mvp import knowledge_engine as e
    records=[{'id':'x','title':'型号XYZ','content':'XYZ是设备型号。','keywords':['XYZ'],'status':'published','visibility':'public','channels':['web']}]
    monkeypatch.setattr(e,'embed',lambda *_args,**_kwargs:(_ for _ in ()).throw(e.EngineError('failed')))
    result,diag=e.hybrid('XYZ是什么？',records,{'x':[1.,0.]},{'rerank_mode':'off'}, {})
    assert result[0]['id']=='x' and diag['embedding']=='keyword_fallback'

def test_rls_covers_all_new_knowledge_tables(make_tenant,make_employee):
    t=make_tenant();a,_=make_employee(t);ready(t,a)
    with transaction() as s:
        for name in ('kb_collections','kb_documents','kb_document_versions','kb_chunks','kb_jobs','kb_engine_configs'):
            assert not s.execute(text('SELECT id FROM '+name)).first()

def test_unsupported_and_oversized_upload_are_explicit(make_tenant,make_employee):
    t=make_tenant();a,_=make_employee(t);collection=a.get('/api/kb/collections').json()[0]['id']
    assert a.post('/api/kb/documents/upload',data={'collection_id':collection},files={'file':('x.exe',b'not executable')}).status_code==422
    assert a.post('/api/kb/documents/upload',content=b'a',headers={'Content-Length':'21000000'}).status_code==413

def test_all_normalized_content_is_chunked_without_legacy_20k_truncation():
    text='产品资料。'*6000+'最后字段XYZ'
    normalized={'title':'资料','blocks':[{'id':'b1','type':'paragraph','text':text,'location':{'page':1}}]}
    pieces=chunks(normalized,max_chars=360)
    assert len(pieces)>20 and '最后字段XYZ' in pieces[-1]['content']
