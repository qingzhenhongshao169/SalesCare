"""Typed, bounded document normalization. Keep original numbers and source locations."""
import csv
import hashlib
import io
import re
import zipfile
from pathlib import Path

FORMATS={'.txt','.md','.markdown','.docx','.pdf','.xlsx','.csv','.pptx'}
MAX_TEXT=2_000_000

class ParseError(Exception):pass

def digest(text):return hashlib.sha256(text.encode('utf-8')).hexdigest()

def decode(data):
    for encoding in (('utf-16',) if data.startswith((b'\xff\xfe',b'\xfe\xff')) else ('utf-8-sig','gb18030')):
        try:return data.decode(encoding)
        except UnicodeError:pass
    raise ParseError('encoding')

def checked_zip(path):
    with zipfile.ZipFile(path) as z:
        files=z.infolist()
        if len(files)>5000 or sum(f.file_size for f in files)>100_000_000:raise ParseError('archive_limit')

def normalize(path,filename):
    path=Path(path);suffix=Path(filename).suffix.lower()
    if suffix not in FORMATS:raise ParseError('unsupported')
    if suffix in ('.docx','.xlsx','.pptx'):checked_zip(path)
    blocks=[];issues=[];title=Path(filename).stem;count=0
    def add(kind,text,location,**extra):
        nonlocal count,title
        text=text.strip()
        if not text:return
        count+=len(text)
        if count>MAX_TEXT:raise ParseError('content_limit')
        if kind=='heading' and len(blocks)==0:title=text[:250]
        blocks.append({'id':'b'+str(len(blocks)+1),'type':kind,'text':text,'location':location,**extra})
    def table(rows,location):
        rows=[[str(c if c is not None else '').strip() for c in row] for row in rows]
        rows=[row for row in rows if any(row)]
        if not rows:return
        if sum(len(row) for row in rows)>50_000:raise ParseError('table_limit')
        add('table','\n'.join(' | '.join(row) for row in rows),location,rows=rows,header=rows[0])
    if suffix in ('.txt','.md','.markdown'):
        text=decode(path.read_bytes())
        paragraph=[];start=1;fenced=False
        def flush():
            if paragraph:add('paragraph','\n'.join(paragraph),{'paragraph':start})
            paragraph.clear()
        for n,line in enumerate(text.splitlines(),1):
            if re.match(r'^\s*(?:```|~~~)',line):fenced=not fenced
            match=re.match(r'^\s{0,3}(#{1,6})\s+(.+)',line) if not fenced else None
            if match:
                flush();add('heading',match[2],{'paragraph':n},level=len(match[1]))
            elif not line.strip():flush()
            else:
                if not paragraph:start=n
                paragraph.append(line)
        flush()
    elif suffix=='.docx':
        from docx import Document
        from docx.text.paragraph import Paragraph
        for n,item in enumerate(Document(path).iter_inner_content(),1):
            if isinstance(item,Paragraph):
                heading=re.search(r'Heading\s*(\d+)',item.style.name)
                add('heading' if heading else 'paragraph',item.text,{'paragraph':n},**({'level':int(heading[1])} if heading else {}))
            else:table([[c.text for c in r.cells] for r in item.rows],{'table':n})
    elif suffix=='.csv':
        text=decode(path.read_bytes())
        try:dialect=csv.Sniffer().sniff(text[:5000])
        except csv.Error:dialect=csv.excel
        table(list(csv.reader(io.StringIO(text),dialect)),{'sheet':'CSV','row_start':1})
    elif suffix=='.xlsx':
        from openpyxl import load_workbook
        book=load_workbook(path,read_only=True,data_only=True)
        formulas=load_workbook(path,read_only=True,data_only=False)
        try:
            for sheet in book:
                if sheet.max_row>10000 or sheet.max_column>200:raise ParseError('table_limit')
                rows=list(sheet.iter_rows(values_only=True))
                for row in formulas[sheet.title].iter_rows():
                    if any(cell.data_type=='f' for cell in row):
                        issues.append({'code':'formula_cache','location':{'sheet':sheet.title},'message':'表格含公式，仅使用已保存的计算值；请核对空值与结果。'})
                        break
                table(rows,{'sheet':sheet.title,'row_start':1})
        finally:book.close();formulas.close()
    elif suffix=='.pptx':
        from pptx import Presentation
        for n,slide in enumerate(Presentation(path).slides,1):
            for shape in slide.shapes:
                if shape.has_table:table([[c.text for c in row.cells] for row in shape.table.rows],{'slide':n})
                elif shape.has_text_frame:add('paragraph',shape.text,{'slide':n})
    elif suffix=='.pdf':
        from pypdf import PdfReader
        import pdfplumber
        reader=PdfReader(path)
        if reader.is_encrypted:raise ParseError('encrypted')
        if len(reader.pages)>300:raise ParseError('page_limit')
        with pdfplumber.open(path) as pdf:
            for n,page in enumerate(pdf.pages,1):
                text=page.extract_text(layout=True) or ''
                if not text.strip():issues.append({'code':'ocr_required','location':{'page':n},'message':'此页没有可读取文字，需要OCR或文字版文件。'})
                else:add('paragraph',text,{'page':n})
                for number,rows in enumerate(page.extract_tables(),1):table(rows,{'page':n,'table':number})
    if not blocks:raise ParseError('ocr_required' if issues else 'empty')
    return {'schema':'salescare.document.v1','title':title,'filename':filename,'format':suffix,
            'blocks':blocks,'issues':issues,'characters':count,'parser_version':'native-v2'}

def chunks(normalized,max_chars=1200):
    result=[];headings=[];pending=[];size=0
    root=normalized['title']
    def path():return [root]+[h['text'] for h in headings if h['text']!=root]
    def identity():return ' · '.join(path())
    def provenance(block):
        return {**block['location'],'heading_path':path()}
    def emit():
        nonlocal pending,size
        if pending:
            text=identity()+'\n'+'\n\n'.join(b['text'] for b in pending)
            result.append({'title':identity()[:250],'content':text,'location':provenance(pending[0]),
                           'block_ids':list(dict.fromkeys([h['id'] for h in headings]+[b['id'] for b in pending])),'digest':digest(text)})
        pending=[];size=0
    for block in normalized['blocks']:
        if block['type']=='heading':
            emit()
            level=block.get('level',1)
            while headings and headings[-1].get('level',1)>=level:headings.pop()
            headings.append(block);continue
        if block['type']=='table':
            emit();rows=block['rows'];header=rows[0];group=[];start=2
            def emit_rows(group,start):
                if not group:return
                text=identity()+'\n'+' | '.join(header)+'\n'+'\n'.join(' | '.join(row) for row in group)
                result.append({'title':identity()[:250],'content':text,'location':{**provenance(block),'row_start':start,'row_end':start+len(group)-1},'block_ids':[h['id'] for h in headings]+[block['id']],'digest':digest(text)})
            for n,row in enumerate(rows[1:],2):
                if group and sum(len(' | '.join(r)) for r in group)+len(' | '.join(row))>max_chars:
                    emit_rows(group,start);group=[];start=n
                group.append(row)
            emit_rows(group,start)
            if len(rows)==1:
                pending=[block];emit()
            continue
        # Preserve full paragraphs when small; split oversized paragraphs with overlap.
        text=block['text']
        if len(text)>max_chars:
            emit()
            for start in range(0,len(text),max_chars-120):
                pending=[{**block,'text':text[start:start+max_chars],'location':{**block['location'],'offset':start}}];emit()
        else:
            if size+len(text)>max_chars:emit()
            pending.append(block);size+=len(text)
    emit()
    if not result:raise ParseError('empty')
    if len(result)>3000:raise ParseError('chunk_limit')
    return result
